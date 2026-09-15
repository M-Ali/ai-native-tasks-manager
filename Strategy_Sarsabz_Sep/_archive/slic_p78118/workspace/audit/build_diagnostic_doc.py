"""Build the Competitive Analysis diagnostic as a Word document.

    uv run --with python-docx --with openpyxl python workspace/audit/build_diagnostic_doc.py

The deck is the argument; this is the evidence behind it. SLIC's evaluation criteria are
"(Qualitative)(Doc Required)" - every scored line needs an uploaded document, and a PPTX
alone is a thin answer to a 15-mark criterion. This produces the readable submission
document: the finding in prose, then the coded tables an evaluator can check.

Date-stamped, never overwritten; a same-day rebuild becomes -v2.
"""

from __future__ import annotations

import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AUDIT = Path(__file__).parent
OUT = AUDIT.parent / "out"
NAVY = RGBColor(0x10, 0x2A, 0x43)
ACCENT = RGBColor(0xC8, 0x10, 0x2E)
MUTED = RGBColor(0x6B, 0x74, 0x80)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def versioned(stem: str, suffix: str) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / f"{stem}_{date.today().isoformat()}{suffix}"
    n = 2
    while p.exists():
        p = OUT / f"{stem}_{date.today().isoformat()}-v{n}{suffix}"
        n += 1
    return p


def para(doc, text="", bold=False, size=10.5, color=None, space=4, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text)
    r.font.size, r.bold, r.italic, r.font.name = Pt(size), bold, italic, "Segoe UI"
    if color:
        r.font.color.rgb = color
    return p


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(it)
        r.font.size, r.font.name = Pt(size), "Segoe UI"


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = NAVY
        r.font.name = "Segoe UI"
    return h


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold, r.font.size, r.font.name = True, Pt(9.5), "Segoe UI"
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size, r.font.name = Pt(9.5), "Segoe UI"
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def load_rows() -> list[dict]:
    wb = openpyxl.load_workbook(AUDIT / "deep" / "capture.xlsx", data_only=True)
    ws = wb["Ads"]
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)
            if r[1].value]


def finding_sections() -> dict[str, str]:
    """Pull the prose straight from the_finding.md so the two never diverge."""
    txt = (AUDIT / "the_finding.md").read_text(encoding="utf-8")
    out, cur, buf = {}, None, []
    for line in txt.split("\n"):
        m = re.match(r"^## (.+)$", line)
        if m:
            if cur:
                out[cur] = "\n".join(buf).strip()
            cur, buf = m.group(1), []
        elif cur:
            buf.append(line)
    if cur:
        out[cur] = "\n".join(buf).strip()
    return out


def clean(md: str) -> list[str]:
    """Markdown prose -> plain paragraphs. Tables are rebuilt natively, so drop them."""
    paras, buf = [], []
    for line in md.split("\n"):
        if line.startswith("|") or line.startswith(">"):
            continue
        if not line.strip():
            if buf:
                paras.append(" ".join(buf))
                buf = []
            continue
        if line.startswith("- ") or line.startswith("1. "):
            if buf:
                paras.append(" ".join(buf))
                buf = []
            paras.append(line)
            continue
        buf.append(line.strip())
    if buf:
        paras.append(" ".join(buf))
    return [re.sub(r"[*`]", "", p) for p in paras if p.strip()]


def main() -> None:
    rows = load_rows()
    n = len(rows)
    sec = finding_sections()

    by_brand = defaultdict(list)
    for r in rows:
        by_brand[r["brand"]].append(r)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Segoe UI", Pt(10.5)

    # ---- cover
    para(doc, "Competitive Analysis", bold=True, size=26, color=NAVY, space=2)
    para(doc, "How Pakistan's life insurance category communicates, and where the gap is",
         size=13, color=MUTED, space=14)
    para(doc, "State Life Insurance Corporation of Pakistan", bold=True, size=11, space=2)
    para(doc, "Prequalification of Advertising Agencies · EPADS reference P78118", space=2)
    para(doc, f"Evaluation criterion: Competitive Analysis (15 marks)", space=2)
    para(doc, f"{n} advertisements · {len(by_brand)} brands · collected 27 August 2026",
         color=MUTED, space=14)
    para(doc, "This document is the evidence behind the accompanying presentation. Every "
              "figure in it is counted from the coded dataset supplied as an annexure, and "
              "every advertisement cited carries a public source URL. Market-structure "
              "figures are sourced to the Competition Commission of Pakistan.",
         italic=True, color=MUTED)
    doc.add_page_break()

    # ---- 1 what we looked at
    heading(doc, "1. What we looked at", 1)
    for p in clean(sec.get("1. What we looked at", "")):
        (bullets(doc, [p[2:]]) if p.startswith("- ") else para(doc, p))
    table(doc, ["Ring", "Brands"], [
        ("Client", "State Life Insurance Corporation of Pakistan"),
        ("Direct competitors", "EFU Life, Jubilee Life, Adamjee Life, IGI Life, Askari Life, TPL Life"),
        ("Adjacent (Takaful)", "Pak-Qatar Family Takaful, Dawood Family Takaful, EFU Life Window Takaful"),
        ("Substitutes", "Al Meezan Investments, UBL Fund Managers"),
        ("No presence", "Postal Life Insurance, Central Directorate of National Savings"),
    ])

    # ---- 2 how the category talks
    heading(doc, "2. How the category talks", 1)
    for p in clean(sec.get("2. How the category talks", "")):
        (bullets(doc, [p[2:]]) if p.startswith("- ") else para(doc, p))

    heading(doc, "Ecosystem convergence", 2)
    para(doc, "Share of each brand's output whose call to action is an owned app or "
              "platform rather than a sale, an agent or a branch.", color=MUTED)
    eco = sorted(((b, sum(1 for r in rs if str(r.get("ecosystem_push")).lower() == "yes"),
                   len(rs)) for b, rs in by_brand.items()), key=lambda t: -t[1])
    table(doc, ["Brand", "Ads driving to an app", "Of", "Share"],
          [(b, k, t, f"{round(100*k/t)}%") for b, k, t in eco if k])

    # ---- 2a why: the mechanism, not just the pattern
    heading(doc, "2a. Why the category is doing this", 1)
    for p2 in clean(sec.get("2a. Why the category is doing this", "")):
        if p2.startswith("- "):
            bullets(doc, [p2[2:]])
        elif p2.startswith("Sources:"):
            para(doc, p2, size=9, color=MUTED, italic=True)
        else:
            para(doc, p2)
    heading(doc, "The structural drivers", 2)
    table(doc, ["Driver", "Evidence"], [
        ("The bank owns the relationship",
         "Private life insurance sells overwhelmingly through bancassurance. The bank "
         "holds the customer, the data and the renewal conversation."),
        ("The bank rations the shelf",
         "CCP 2025: banks “impose additional internal limits on the amount of business "
         "insurance companies can conduct through banks” - characterised as a refusal to deal."),
        ("The channel carries a trust problem",
         "CCP 2025: bank and insurer staff “do not properly guide the customers”; "
         "terms presented in fine print."),
        ("Almost nobody has bought the product",
         "3% of Pakistanis hold a life policy. Penetration 0.67% of GDP against 6.7% "
         "globally; density US$14 per head against India's US$82."),
    ])

    # ---- 2b audience
    heading(doc, "2b. Who the category is talking to", 1)
    gens = Counter(r.get("audience_generation") for r in rows)
    labels = [("geny", "Millennial (~28-43)"), ("mixed", "Multi-generational"),
              ("genx", "Gen X (~44-59)"), ("genz", "Gen Z (under ~28)"),
              ("boomer_plus", "60+"), ("unclear", "Unclear")]
    table(doc, ["Cast for", "Ads", "Share"],
          [(lab, gens.get(k, 0), f"{round(100*gens.get(k, 0)/n)}%")
           for k, lab in labels if gens.get(k)])
    for p2 in clean(sec.get("2b. Who the category is talking to", "")):
        (bullets(doc, [p2[2:]]) if p2.startswith("- ") else para(doc, p2))

    # ---- 2c three readings
    heading(doc, "2c. Three readings that change the brief", 1)
    for p2 in clean(sec.get("2c. Three readings that change the brief", "")):
        (bullets(doc, [p2[2:]]) if p2.startswith("- ") else para(doc, p2))
    heading(doc, "Calendar content as a share of output", 2)
    cal = sorted(((b, sum(1 for r in rs if r.get("content_type") == "calendar_topical"),
                   len(rs)) for b, rs in by_brand.items()), key=lambda t: -t[1])
    table(doc, ["Brand", "Calendar posts", "Of", "Share"],
          [(b, k, t, f"{round(100*k/t)}%") for b, k, t in cal])
    heading(doc, "Production investment and who appears in the work", 2)
    table(doc, ["Brand", "Studio shoots", "Customer in frame", "Executive / staff in frame"],
          [(b,
            sum(1 for r in rs if r.get("production_value") == "studio"),
            sum(1 for r in rs if r.get("who_is_in_frame") == "customer"),
            sum(1 for r in rs if r.get("who_is_in_frame") in ("executive", "staff")))
           for b, rs in sorted(by_brand.items(),
                               key=lambda kv: -sum(1 for r in kv[1]
                                                   if r.get("production_value") == "studio"))])

    # ---- 3 proof
    heading(doc, "3. What nobody is saying", 1)
    for p in clean(sec.get("3. What nobody is saying", "")):
        (bullets(doc, [p[2:]]) if p.startswith("- ") else para(doc, p))
    pf = Counter(r.get("proof_device") or "none" for r in rows)
    table(doc, ["Proof device", "Ads", "Share"],
          [(k, v, f"{round(100*v/n)}%") for k, v in pf.most_common()])

    heading(doc, "The three performance numbers published in this category", 2)
    table(doc, ["Advertisement", "Brand", "Number"], [
        ("E-Kachehri, 13 August 2026", "State Life", "745 complaints received, 744 resolved"),
        ("E-Kachehri, 22 July 2026", "State Life", "730 received, 729 resolved"),
        ("Claims paid, 29 June 2026", "Jubilee Life", "PKR 54 billion paid in 2025"),
    ])

    # ---- 4 where brands stand
    heading(doc, "4. Where the brands actually stand", 1)
    for p in clean(sec.get("4. Where the brands actually stand", "")):
        (bullets(doc, [p[2:]]) if p.startswith("- ") else para(doc, p))
    heading(doc, "Content mix by brand", 2)
    types = ["brand_building", "tactical_promo", "product_feature", "corporate_pr",
             "calendar_topical", "csr", "recruitment", "ugc_repost"]
    present = [t for t in types if any(r.get("content_type") == t for r in rows)]
    table(doc, ["Brand"] + [t.replace("_", " ") for t in present] + ["Building %"],
          [[b] + [sum(1 for r in rs if r.get("content_type") == t) or "-" for t in present]
           + [f"{round(100*sum(1 for r in rs if r.get('content_type')=='brand_building')/len(rs))}%"]
           for b, rs in sorted(by_brand.items(),
                               key=lambda kv: -sum(1 for r in kv[1]
                                                   if r.get("content_type") == "brand_building"))])

    # ---- 5 what State Life can own
    heading(doc, "5. What State Life can credibly own", 1)
    for p in clean(sec.get("5. What State Life can own", "")):
        if p.startswith("- "):
            bullets(doc, [p[2:]])
        elif re.match(r"^\d\. ", p):
            bullets(doc, [p[3:]])
        else:
            para(doc, p)

    # ---- 6 the choice
    heading(doc, "6. The choice to put to the client", 1)
    for p in clean(sec.get("6. The choice to put to the client", "")):
        (bullets(doc, [p[2:]]) if p.startswith("- ") else para(doc, p))

    # ---- 7 method and limits
    doc.add_page_break()
    heading(doc, "7. Method, and what this evidence cannot show", 1)
    para(doc, "Stating the limits is what separates an audit from an assertion. Each of "
              "these is a property of the available data, not an omission.")
    bullets(doc, [
        "Share of voice is not measurable here. Every brand contributed the same number of "
        "posts because that is the platform's logged-out ceiling, so relative volume is an "
        "artefact of collection. No share-of-voice figure appears in this document.",
        "All material is organic owned content. It carries no information about media "
        "spend, reach or weight. The Meta Ad Library is the only public source for paid "
        "creative and publishes no spend for commercial advertisers.",
        f"{sum(1 for r in rows if r.get('coder') == 'claude-brand-default')} of {n} posts "
        "carry no caption and are coded to a brand-typical default rather than an "
        "individual reading. They are labelled in the dataset.",
        "Audience-generation codes are inference from casting, styling and language. They "
        "describe who an advertisement addresses, never who buys the product.",
        "Coding is per advertisement against a fixed nine-dimension schema, supplied as an "
        "annexure so any figure here can be recounted.",
    ])

    heading(doc, "Coding schema", 2)
    table(doc, ["Dimension", "What it records"], [
        ("product_focus", "The product, platform or ecosystem the advertisement pushes"),
        ("audience_generation", "Who the work is cast for (inference)"),
        ("content_type", "Brand-building, tactical, product, PR, calendar, CSR, recruitment"),
        ("claim_primary / secondary", "What the advertisement claims"),
        ("register", "Emotional key: fear, aspiration, duty, reassurance, pride, humour"),
        ("proof_device", "What substantiates the claim, or none"),
        ("who_is_in_frame", "Customer, celebrity, executive, staff, expert, no people"),
        ("production_value", "Studio, stock, template, event photography, UGC, AI"),
        ("ecosystem_push", "Whether the call to action is an owned app or platform"),
    ])

    heading(doc, "Sample frame", 2)
    table(doc, ["Brand", "Ring", "Ads coded", "Followers"],
          [(b, (rs[0].get("ring") or "-"), len(rs), "") for b, rs in by_brand.items()])

    path = versioned("SLIC_P78118_competitive_analysis", ".docx")
    doc.save(path)
    print(f"{path}")
    print(f"  {n} ads, {len(by_brand)} brands, {len(doc.paragraphs)} paragraphs, "
          f"{len(doc.tables)} tables")


if __name__ == "__main__":
    main()
