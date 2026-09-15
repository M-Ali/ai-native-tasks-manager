"""Deliverables: submission checklist (Markdown + Excel) and the Annexure forms (Word)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document as Docx
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import criteria
from .readiness import Assessment

OUT_DIR = Path("workspace/out")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILL = {
    "ready": PatternFill("solid", fgColor="C6E0B4"),
    "partial": PatternFill("solid", fgColor="FFE699"),
    "missing": PatternFill("solid", fgColor="F8CBAD"),
    "na": PatternFill("solid", fgColor="E7E6E6"),
}


def _stamp() -> str:
    return date.today().isoformat()


def _versioned(out_dir: Path, stem: str, suffix: str) -> Path:
    """Never overwrite a prior deliverable: date-stamp, then add -v2, -v3 as needed."""
    out_dir.mkdir(parents=True, exist_ok=True)
    base = f"{stem}_{_stamp()}"
    path = out_dir / f"{base}{suffix}"
    n = 2
    while path.exists():
        path = out_dir / f"{base}-v{n}{suffix}"
        n += 1
    return path


# --------------------------------------------------------------------------- markdown


def checklist_markdown(assessment: Assessment, out_dir: Path = OUT_DIR) -> Path:
    t = criteria.Tender()
    a = assessment
    lines: list[str] = []
    add = lines.append

    add(f"# SLIC Prequalification {t.reference} - submission checklist")
    add("")
    add(f"**Agency:** {a.agency or '(unnamed)'}  ")
    add(f"**Procuring agency:** {t.agency}  ")
    add(f"**Submission deadline:** {t.submission_deadline} (EPADS v2.0, opening {t.opening})  ")
    add(f"**Clarification date:** {t.clarification_deadline}  ")
    add(f"**Selection:** {t.selection} - {t.total_marks} technical marks, {t.passing_marks} to pass  ")
    add(f"**Generated:** {_stamp()}")
    add("")

    verdict = "PASSES" if a.passes else "BELOW PASS MARK"
    add(f"## Score as assessed: {a.earned:g} / {a.available} -- {verdict}")
    add("")
    add("| Bucket | Earned | Available |")
    add("|---|---:|---:|")
    for bucket, (got, avail) in a.by_bucket().items():
        add(f"| {criteria.BUCKET_LABELS.get(bucket, bucket)} | {got:g} | {avail} |")
    add(f"| **Total** | **{a.earned:g}** | **{a.available}** |")
    add("")

    add("## 1. Eligibility documents (pass/fail, PQ doc p.21-22)")
    add("")
    add("| # | Requirement | Status | Note |")
    add("|---:|---|---|---|")
    for n, ln in enumerate(a.eligibility, 1):
        add(f"| {n} | {ln.label} | {ln.status} | {ln.note} |")
    add("")

    add("## 2. Technical evaluation (PQ doc p.22-23)")
    add("")
    scored = criteria.scored_by_key()
    for bucket, label in criteria.BUCKET_LABELS.items():
        rows = [ln for ln in a.lines if ln.bucket == bucket]
        if not rows:
            continue
        got = sum(ln.earned for ln in rows)
        avail = sum(ln.marks for ln in rows)
        add(f"### {label} -- {got:g}/{avail}")
        add("")
        add("| Criterion | Marks | Status | Earned | Evidence required |")
        add("|---|---:|---|---:|---|")
        for ln in rows:
            add(
                f"| {ln.label} | {ln.marks} | {ln.status} | {ln.earned:g} | "
                f"{scored[ln.key].evidence} |"
            )
        add("")

    gaps = a.gaps()
    if gaps:
        add("## 3. Where the marks are being lost")
        add("")
        add("| Criterion | Marks at stake | Status | Next action |")
        add("|---|---:|---|---|")
        for ln in gaps:
            action = ln.note or scored[ln.key].evidence
            add(f"| {ln.label} | {ln.marks - ln.earned:g} | {ln.status} | {action} |")
        add("")

    eg = a.eligibility_gaps()
    if eg:
        add("## 4. Eligibility items still outstanding")
        add("")
        for ln in eg:
            add(f"- **{ln.label}** - {ln.status}{(' - ' + ln.note) if ln.note else ''}")
        add("")

    path = _versioned(out_dir, "SLIC_P78118_checklist", ".md")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# ------------------------------------------------------------------------------ excel


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def checklist_excel(assessment: Assessment, out_dir: Path = OUT_DIR) -> Path:
    t = criteria.Tender()
    a = assessment
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Tender reference", t.reference),
        ("Title", t.title),
        ("Procuring agency", t.agency),
        ("Portal", t.portal),
        ("Method", t.method),
        ("Selection", t.selection),
        ("Clarification date", t.clarification_deadline),
        ("Submission deadline", t.submission_deadline),
        ("Opening", t.opening),
        ("Passing marks", f"{t.passing_marks} / {t.total_marks}"),
        ("Agency", a.agency),
        ("Score as assessed", f"{a.earned:g} / {a.available}"),
        ("Verdict", "PASSES" if a.passes else "BELOW PASS MARK"),
        ("Generated", _stamp()),
    ]
    for r, (k, v) in enumerate(rows, 1):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
    _autosize(ws, {1: 26, 2: 78})

    ws = wb.create_sheet("Eligibility")
    headers = ["#", "Requirement", "Status", "Note", "PQ page"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    pages = {i.key: i.page for i in criteria.ELIGIBILITY}
    for r, ln in enumerate(a.eligibility, 2):
        ws.cell(r, 1, r - 1)
        ws.cell(r, 2, ln.label).alignment = Alignment(wrap_text=True, vertical="top")
        sc = ws.cell(r, 3, ln.status)
        sc.fill = STATUS_FILL.get(ln.status, STATUS_FILL["na"])
        ws.cell(r, 4, ln.note).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 5, pages[ln.key])
    ws.freeze_panes = "A2"
    _autosize(ws, {1: 5, 2: 70, 3: 12, 4: 40, 5: 9})

    ws = wb.create_sheet("Technical scoring")
    headers = ["Bucket", "Criterion", "Marks", "Status", "Earned", "Evidence required", "Note"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    scored = criteria.scored_by_key()
    r = 2
    for ln in a.lines:
        ws.cell(r, 1, criteria.BUCKET_LABELS.get(ln.bucket, ln.bucket))
        ws.cell(r, 2, ln.label).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 3, ln.marks)
        sc = ws.cell(r, 4, ln.status)
        sc.fill = STATUS_FILL.get(ln.status, STATUS_FILL["na"])
        ws.cell(r, 5, ln.earned)
        ws.cell(r, 6, scored[ln.key].evidence).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        ws.cell(r, 7, ln.note).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.cell(r, 2, "TOTAL").font = Font(bold=True)
    ws.cell(r, 3, a.available).font = Font(bold=True)
    ws.cell(r, 5, a.earned).font = Font(bold=True)
    ws.freeze_panes = "A2"
    _autosize(ws, {1: 30, 2: 60, 3: 8, 4: 12, 5: 9, 6: 50, 7: 34})

    ws = wb.create_sheet("Gaps")
    headers = ["Criterion", "Marks at stake", "Status", "Next action"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r, ln in enumerate(a.gaps(), 2):
        ws.cell(r, 1, ln.label).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 2, ln.marks - ln.earned)
        ws.cell(r, 3, ln.status)
        ws.cell(r, 4, ln.note or scored[ln.key].evidence).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    ws.freeze_panes = "A2"
    _autosize(ws, {1: 60, 2: 15, 3: 12, 4: 55})

    path = _versioned(out_dir, "SLIC_P78118_checklist", ".xlsx")
    wb.save(path)
    return path


# ------------------------------------------------------------------------------- word


def _para(doc, text: str, bold: bool = False, size: int = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def annexure_forms(agency: str, out_dir: Path = OUT_DIR) -> Path:
    """Annexure I (letter of application), II (declaration) and III (staff table)."""
    doc = Docx()
    name = agency or "[Name of Agency]"

    doc.add_heading("Annexure I - Letter of Application", level=1)
    _para(doc, "To:\nDivisional Head (CAD)\nState Life Insurance Corporation of Pakistan,\n"
               "Principal Office, Dr Ziauddin Ahmed Road, Karachi.")
    _para(doc, "Respected Madam,")
    _para(
        doc,
        f'Being duly authorized to represent and act on behalf of {name} (hereinafter '
        f'"the Agency") and having reviewed and fully understood all the prequalification '
        f"information provided, the undersigned hereby apply to be prequalified as Agency "
        f"for providing Advertising Services to SLIC.",
    )
    _para(
        doc,
        "Attached to this letter are attested true copies (of original documents) as "
        "required per the Evaluation Criteria. SLIC and its authorized representatives are "
        "hereby authorized to conduct any inquiries or investigations to verify the "
        "statements, documents and information submitted with this application, and to seek "
        "clarification from our bankers and clients on any financial and technical aspect.",
    )
    _para(
        doc,
        "This application is made with the full understanding that documents by prequalified "
        "agencies will be subject to verification, and that SLIC reserves the right to amend "
        "the scope of this project or to cancel the prequalification process and reject "
        "applications in accordance with the Public Procurement Rules.",
    )
    _para(doc, "\nName & Designation: ______________________")
    _para(doc, f"For and on behalf of: {name}")
    _para(doc, "Agency stamp to be affixed")

    doc.add_page_break()
    doc.add_heading("Annexure II - To Whom It May Concern", level=1)
    _para(doc, "(To be executed on a Rs. 100/- stamp paper)", bold=True)
    _para(
        doc,
        f"This is to certify that {name}, having its registered office at [Address], hereby "
        "confirms and declares that:",
    )
    for text in (
        "The Agency has never been involved in any criminal, unlawful or illegal activities.",
        "The Agency has not been blacklisted by any entity, organization or regulatory "
        "authority, including but not limited to APNS, PBA, or any other relevant body.",
        "The Agency undertakes that all information and documents submitted to SLIC for "
        "prequalification are genuine, accurate and free from any falsification or forgery.",
        "The Agency acknowledges that if any information or document submitted is found to "
        "be false, forged or misleading at any stage, the Agency shall be disqualified "
        "immediately and SLIC reserves the right to take further legal action.",
        "The Agency declares that there is no existing or potential conflict of interest "
        "affecting its ability to perform the required services impartially, and commits to "
        "disclose any such conflict to SLIC immediately should it arise.",
        "We further affirm that all information provided is true and accurate to the best of "
        "our knowledge.",
    ):
        doc.add_paragraph(text, style="List Number")
    for label in (
        "Authorized Signatory", "Name", "Designation", "CNIC", "Agency Seal", "Date",
    ):
        _para(doc, f"{label}: ______________________________")

    doc.add_page_break()
    doc.add_heading("Annexure III - Particulars of Staff", level=1)
    for text in (
        "Provide details of all functional and support staff in the format below.",
        "Enclose CVs of each listed staff member with relevant experience and role.",
        "Clearly mention employment status (permanent or contractual).",
        "Ensure the information is accurate and verifiable; false details may lead to "
        "disqualification.",
    ):
        doc.add_paragraph(text, style="List Number")
    headers = [
        "S.No", "Name", "Designation", "Qualification", "Total Exp (Years)",
        "Key Expertise/Specialization", "Employment status", "Date of joining", "CV enclosed",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for c, h in enumerate(headers):
        cell = table.rows[0].cells[c]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for _ in range(12):
        table.add_row()

    path = _versioned(out_dir, "SLIC_P78118_annexures", ".docx")
    doc.save(path)
    return path
