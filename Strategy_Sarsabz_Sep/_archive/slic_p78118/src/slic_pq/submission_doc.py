"""Render a Markdown submission section as a Word document.

Every scored criterion in this tender is `(Qualitative)(Doc Required)` — each one needs an
uploaded document, so each written section has to become a .docx. Keeping the Markdown as
the single source and rendering from it means the prose cannot drift between the working
file and the thing that gets submitted.

Markdown tables become real Word tables; inline emphasis survives as bold and italic runs.

Two rendering rules exist because both were shipped as bugs first:

- **List items and blockquotes wrap across lines in the source.** Their continuations must
  be joined before rendering, or emphasis spanning the break loses half its markers and
  leaves a stray asterisk on the page.
- **Nested emphasis is not supported**, because it is ambiguous in Markdown itself. Write
  ``**BOLD** — *italic*`` rather than ``**BOLD — *italic***``.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x10, 0x2A, 0x43)
ACCENT = RGBColor(0xC8, 0x10, 0x2E)

_EMPHASIS = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
_ORDERED = re.compile(r"^\d+\. ")


def _versioned(out_dir: Path, stem: str, suffix: str = ".docx") -> Path:
    """Date-stamp, never overwrite. A same-day rebuild becomes -v2, -v3."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{stem}_{date.today().isoformat()}{suffix}"
    n = 2
    while path.exists():
        path = out_dir / f"{stem}_{date.today().isoformat()}-v{n}{suffix}"
        n += 1
    return path


def _rich(par, text: str, size: float = 10.5, bold: bool = False,
          color: RGBColor | None = None) -> None:
    """Render inline emphasis as runs rather than leaving markers on the page."""
    for chunk in _EMPHASIS.split(text):
        if not chunk:
            continue
        is_bold, is_italic = bold, False
        if chunk.startswith("**") and chunk.endswith("**"):
            chunk, is_bold = chunk[2:-2], True
        elif chunk.startswith("*") and chunk.endswith("*"):
            chunk, is_italic = chunk[1:-1], True
        elif chunk.startswith("`") and chunk.endswith("`"):
            chunk = chunk[1:-1]
        run = par.add_run(chunk)
        run.bold, run.italic = is_bold, is_italic
        run.font.size, run.font.name = Pt(size), "Segoe UI"
        if color is not None:
            run.font.color.rgb = color


def _is_table_head(lines: list[str], i: int) -> bool:
    return (lines[i].startswith("|") and i + 1 < len(lines)
            and set(lines[i + 1].replace("|", "").strip()) <= set("-: ")
            and lines[i + 1].strip() != "")


def render(markdown_path: Path, out_dir: Path, stem: str) -> Path:
    """Render `markdown_path` to a date-stamped .docx in `out_dir`. Returns the path."""
    lines = markdown_path.read_text(encoding="utf-8").split("\n")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Segoe UI", Pt(10.5)

    i, first_title = 0, True
    while i < len(lines):
        line = lines[i].rstrip()

        if _is_table_head(lines, i):
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            body = []
            while i < len(lines) and lines[i].startswith("|"):
                body.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style, table.alignment = "Light Grid Accent 1", WD_TABLE_ALIGNMENT.LEFT
            for j, head in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = ""
                _rich(cell.paragraphs[0], head, 9, bold=True)
            for row in body:
                cells = table.add_row().cells
                for j, value in enumerate(row[:len(header)]):
                    cells[j].text = ""
                    _rich(cells[j].paragraphs[0], value, 9)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if line.startswith("# "):
            if first_title:
                _rich(doc.add_paragraph(), line[2:], 26, bold=True, color=NAVY)
                first_title = False
            else:
                _rich(doc.add_heading("", 1), line[2:], 16, bold=True, color=NAVY)
        elif line.startswith("### "):
            _rich(doc.add_heading("", 2), line[4:], 13, bold=True, color=ACCENT)
        elif line.startswith("## "):
            _rich(doc.add_heading("", 1), line[3:], 16, bold=True, color=NAVY)
        elif line.startswith("---"):
            doc.add_page_break()
        elif line.startswith("- ") or _ORDERED.match(line):
            bullet = line.startswith("- ")
            buf = [line[2:] if bullet else _ORDERED.sub("", line)]
            while (i + 1 < len(lines) and lines[i + 1].startswith("  ")
                   and lines[i + 1].strip()
                   and not lines[i + 1].lstrip().startswith(("- ", "|", "#"))):
                i += 1
                buf.append(lines[i].strip())
            par = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            par.paragraph_format.space_after = Pt(3)
            _rich(par, " ".join(buf))
        elif line.startswith("> "):
            buf = [line[2:]]
            while i + 1 < len(lines) and lines[i + 1].startswith(">"):
                i += 1
                buf.append(lines[i].lstrip("> ").rstrip())
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            par.paragraph_format.space_after = Pt(6)
            _rich(par, " ".join(x for x in buf if x), bold=True, color=NAVY)
        elif line.strip():
            buf = [line.strip()]
            while (i + 1 < len(lines) and lines[i + 1].strip()
                   and not lines[i + 1].startswith(("#", "|", "- ", "> ", "---"))
                   and not _ORDERED.match(lines[i + 1])):
                i += 1
                buf.append(lines[i].strip())
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(6)
            _rich(par, " ".join(buf))
        i += 1

    path = _versioned(out_dir, stem)
    doc.save(path)
    return path
