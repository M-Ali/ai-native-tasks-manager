"""Render the analysis Markdown to .docx, tables included.

    python render_docx.py <analysis.md> --out <dir> [--stem NAME]

Markdown is the source and this is only a render, so the working file and the delivered
file cannot drift. Never hand-edit the .docx: the next render silently discards it.

Supports the subset the analysis actually uses - headings, bullets, numbered lists,
pipe tables, blockquotes, **bold**/*italic* inline, and `> ` callouts. Anything else is
written as plain text rather than dropped, because losing a paragraph silently from a
client document is worse than losing its formatting.

Output is date-stamped and versioned (-v2, -v3 on the same day) rather than overwritten.
A previous version of an analysis is evidence of what was said and when.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

RE_BOLD = re.compile(r"\*\*(.+?)\*\*")
RE_ITAL = re.compile(r"(?<!\*)\*([^*]+?)\*(?!\*)")
RE_CODE = re.compile(r"`([^`]+?)`")


def add_runs(par, text: str) -> None:
    """Inline **bold**, *italic* and `code`, applied without nesting."""
    tokens = re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if m := RE_BOLD.fullmatch(tok):
            par.add_run(m.group(1)).bold = True
        elif m := RE_CODE.fullmatch(tok):
            r = par.add_run(m.group(1))
            r.font.name = "Consolas"
        elif m := RE_ITAL.fullmatch(tok):
            par.add_run(m.group(1)).italic = True
        else:
            par.add_run(tok)


def is_table_sep(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?[\s:|-]+\|[\s:|-]*", line)) and "-" in line


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--stem", default=None)
    args = ap.parse_args()

    import docx
    from docx.shared import Pt

    if not args.source.exists():
        sys.exit(f"{args.source} not found")
    lines = args.source.read_text(encoding="utf-8").splitlines()
    args.out.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # table: header, separator, body
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            header = split_row(stripped)
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(split_row(lines[i].strip()))
                i += 1
            width = max([len(header)] + [len(r) for r in body])
            t = doc.add_table(rows=1, cols=width)
            t.style = "Light Grid Accent 1"
            for j, cell in enumerate(header):
                p = t.rows[0].cells[j].paragraphs[0]
                add_runs(p, cell)
                for r in p.runs:
                    r.bold = True
            for row in body:
                cells = t.add_row().cells
                for j in range(width):
                    add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else "")
            doc.add_paragraph()
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            if level == 1:
                h = doc.add_heading("", level=0)
            else:
                h = doc.add_heading("", level=min(level - 1, 4))
            add_runs(h, text)
        elif stripped.startswith(("- ", "* ")):
            add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+[.)]\s", stripped):
            add_runs(doc.add_paragraph(style="List Number"),
                     re.sub(r"^\d+[.)]\s+", "", stripped))
        elif stripped.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, stripped.lstrip(">").strip())
        elif set(stripped) <= {"-", "*", "_"} and len(stripped) >= 3:
            doc.add_paragraph()
        else:
            add_runs(doc.add_paragraph(), stripped)
        i += 1

    stem = args.stem or args.source.stem
    today = dt.date.today().isoformat()
    path = args.out / f"{stem}_{today}.docx"
    n = 2
    while path.exists():
        path = args.out / f"{stem}_{today}-v{n}.docx"
        n += 1
    doc.save(path)
    print(f"{path}")
    print("Markdown is the source. Edit the .md and re-render; never edit the .docx.")


if __name__ == "__main__":
    main()
