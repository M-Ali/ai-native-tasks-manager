# -*- coding: utf-8 -*-
"""Write hand-coded reads for all 84 posts into scan/capture.xlsx (Ads sheet).

Claim vocabulary REWRITTEN for facial skincare (Pakistan), per SKILL.md - the shipped
list is tuned for financial services. Every code below was read off the contact sheets,
image by image; `unclear` is used where the thumbnail could not be read.
"""
import csv
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

WS = Path(r"D:/Personal/SL_aug2026/.claude/skills/category-creative-scan-workspace/iteration-1/brand-wise-deck-request/with_skill/scan")
CODER = "claude-scan"

# shortcode: (product_focus, gen, content_type, claim1, claim2, register, proof, frame,
#             production, eco_push, campaign_platform, addresses_objection, lang, notes)
C = {
# ---------------- POND'S Pakistan ----------------
"DZHdiEyhZDy": ("brand", "geny", "brand_building", "brand_corporate", "", "aspiration", "celebrity", "celebrity", "studio", "no", "", "no", "english_first", "'BRAND AMBASSADOR' reveal on a pink studio set; silhouette, no product."),
"Db3A2sdAM4I": ("brand", "geny", "brand_building", "other", "", "aspiration", "none", "customer", "ugc", "no", "", "no", "other", "Woman in hijab, domestic interior. No product or copy legible at thumbnail; coded from casting only."),
"Db8h3P8NH_Q": ("brand", "genz", "product_feature", "ingredient_science", "", "reassurance", "expert_authority", "customer", "ugc", "no", "", "yes", "english_first", "In-store reel: 'The girls and their skin concerns!! ft. skincare specialist'."),
"DbdU1hRM3RC": ("Bright Beauty (Niacinamide)", "unclear", "product_feature", "brightening_tone", "ingredient_science", "aspiration", "claim_statistics", "none_people", "studio", "no", "", "yes", "english_first", "Pack shot: 'NIACINAMIDE [VITAMIN B3] ... 4X' on a towel."),
"DbddJl7xO_h": ("Bright Beauty (Niacinamide)", "genz", "product_feature", "brightening_tone", "", "aspiration", "none", "customer", "ugc", "no", "", "no", "english_first", "Hand holding the pink jar, caption 'under 15 seconds'."),
"DbnoOIGMTP-": ("POND'S Community", "genz", "tactical_promo", "sun_protection", "", "humour", "none", "none_people", "event_photo", "no", "PONDS Community", "no", "english_first", "'Sunscreen Trivia with the PONDS Community...' balloons, office photo."),
"DbnpzicDZOn": ("POND'S Community", "genz", "tactical_promo", "other", "", "humour", "none", "none_people", "template_graphic", "yes", "PONDS Community", "no", "english_first", "'GIVEAWAY WINNER ... Join our Broadcast channel for the reveal!' - drives to an owned IG broadcast channel."),
"DbtGImXpDl_": ("brand", "genz", "product_feature", "ingredient_science", "", "reassurance", "expert_authority", "customer", "ugc", "no", "", "yes", "english_first", "Second cut of the in-store skincare-specialist reel."),
"DbvTihDDQH0": ("Sun Miracle De-Tan", "geny", "product_feature", "sun_protection", "brightening_tone", "aspiration", "none", "customer", "studio", "no", "", "yes", "english_first", "'POND'S SUN MIRACLE DETAN Face Wash' held to camera by a model."),
"DcOZTuIjVnU": ("brand", "unclear", "brand_building", "other", "", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "'YOUR SKIN'S LATEST OBSESSION' over cream texture - teaser, no product named."),
"DcQ6QiKNIV-": ("Sun Miracle", "unclear", "brand_building", "sun_protection", "", "humour", "none", "none_people", "studio", "no", "", "no", "english_first", "'SOAK UP THE SUN, NOT THE TAN' - tubes pegged on a washing line against sky. The one art-directed idea in the set."),
"Dca-DsttT4z": ("Bright Beauty (Niacinamide)", "genz", "product_feature", "brightening_tone", "", "aspiration", "none", "customer", "ugc", "no", "", "no", "english_first", "Unboxing collage, 'the skincare'."),
# ---------------- Glow & Lovely Pakistan ----------------
"Db2bAlqijPb": ("Bright C Glow Facewash", "genz", "tactical_promo", "price_value", "", "aspiration", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "Quiz sticker: 'Ever heard of a facewash that gives you glow for just Rs. 99?'"),
"Dbde5OOuRbu": ("Re-New Bright", "geny", "brand_building", "brightening_tone", "", "aspiration", "none", "none_people", "studio", "no", "Renew Bright", "no", "english_first", "'Wherever summer takes you, Renew Bright's already there.' Pink suitcase."),
"Dcik1OpC0g0": ("Bright C Glow Facewash", "genz", "tactical_promo", "brightening_tone", "", "humour", "none", "none_people", "template_graphic", "no", "", "no", "urdu_first", "'AB NAHI' headline with a BUY NOW button."),
"DcLm2vCibio": ("Re-New Bright", "genz", "brand_building", "brightening_tone", "", "humour", "none", "none_people", "template_graphic", "no", "Renew Bright", "no", "english_first", "Mocked-up ChatGPT 5.1 screenshot: 'Make this brighter' / 'can't go brighter than renew bright'."),
"DcTbukOl9Lx": ("Bright C Glow Facewash", "genz", "product_feature", "cleansing_purity", "", "aspiration", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "'HOLD & SCROLL FOR THE SECRET BEHIND FRESHER SKIN'."),
"DcTciSnCloR": ("Bright C Glow Facewash", "geny", "product_feature", "ingredient_science", "brightening_tone", "reassurance", "none", "customer", "stock", "no", "", "yes", "english_first", "'BRIGHT C FACEWASH ... KEY INGREDIENTS: VITAMIN C + PURE LEMON'; model reads as licensed library imagery."),
"DcVODRhlxnP": ("brand", "genz", "brand_building", "other", "", "aspiration", "none", "none_people", "stock", "no", "", "no", "english_first", "'WHAT'S INSIDE OUR MAKEUP BAG!' - lifestyle engagement post carrying no skincare claim."),
"DcWmISVin8z": ("Re-New Bright", "genz", "product_feature", "brightening_tone", "", "aspiration", "claim_statistics", "none_people", "template_graphic", "no", "Renew Bright", "yes", "english_first", "'MY LOVE LANGUAGE IS MAKING YOUR SKIN 98% BRIGHTER'."),
"DcX4kOGuL3S": ("Re-New Bright", "unclear", "product_feature", "ingredient_science", "brightening_tone", "reassurance", "none", "none_people", "template_graphic", "no", "Renew Bright", "yes", "english_first", "'A RESET FOR YOUR SKIN - helps reduce dullness / Vitamin C + Niacinamide power / evens skin tone'."),
"DcXtE0juvss": ("Bright C Glow Facewash", "unclear", "product_feature", "brightening_tone", "", "aspiration", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "'BRIGHTER LOOKING SKIN! NEW Glow & Lovely BRIGHT C GLOW FACEWASH'."),
"DcYcmJ_OADZ": ("Bright C Glow Facewash", "unclear", "product_feature", "natural_ingredients", "brightening_tone", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "Lemon slices with the tube."),
"DcZuVcUChB0": ("Re-New Bright", "geny", "product_feature", "brightening_tone", "", "aspiration", "none", "customer", "studio", "no", "Renew Bright", "no", "english_first", "'REMINDER ... NEVER SKIP RENEW BRIGHT, YOUR SKIN WILL THANK YOU LATER!'"),
# ---------------- Garnier Pakistan ----------------
"Da0DBhVj4H3": ("Micellar Cleansing Oil-in-Water", "unclear", "product_feature", "cleansing_purity", "", "reassurance", "none", "none_people", "studio", "no", "", "no", "english_first", "Petri-dish split: 'OIL BASED CLEANSER / WATER INFUSED CLEANSER'."),
"Da5JDWUDlGz": ("Micellar range", "genz", "product_feature", "cleansing_purity", "", "humour", "none", "none_people", "template_graphic", "no", "Choose your fighter", "yes", "english_first", "'CHOOSE YOUR FIGHTER - Micellar All in 1 VS Micellar Oil in Water'."),
"DaiQEIeqndM": ("Micellar Cleansing Oil-in-Water", "genz", "ugc_repost", "cleansing_purity", "", "reassurance", "testimonial", "customer", "ugc", "no", "", "yes", "english_first", "Repost @wania nadeem: 'NO RESIDUE / NO DRYNESS'."),
"DaiZhLZm7MK": ("Micellar Cleansing Oil-in-Water", "unclear", "product_feature", "cleansing_purity", "", "reassurance", "none", "none_people", "studio", "no", "", "no", "english_first", "Bottle on a yellow ground."),
"Daid9aDsk9H": ("Vitamin C serum", "genz", "ugc_repost", "brightening_tone", "", "aspiration", "testimonial", "customer", "ugc", "no", "", "yes", "english_first", "Repost @bakhtawar laiba: 'IT REDUCES...' serum held to camera."),
"DapqDhFACp2": ("Vitamin C serum", "genz", "brand_building", "brightening_tone", "", "humour", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "Line illustration of a figure dragging a serum bottle: 'C'MON, BRIGHTEN MY SKIN'."),
"DavOHQ7g_yu": ("Micellar Cleansing Oil-in-Water", "geny", "ugc_repost", "cleansing_purity", "", "aspiration", "celebrity", "celebrity", "ugc", "no", "", "no", "english_first", "Repost @Merub Ali (actor)."),
"DbLSxUuIpwE": ("Micellar Cleansing Oil-in-Water", "unclear", "product_feature", "cleansing_purity", "", "reassurance", "none", "none_people", "studio", "no", "", "yes", "english_first", "Benefit bubbles: 'ZERO RESIDUE / HYDRATES & SOOTHES SKIN / NO RINSE NEEDED / REMOVES IMPURITIES & DIRT'."),
"DbTez7ZNSPV": ("Micellar Cleansing Oil-in-Water", "genz", "ugc_repost", "cleansing_purity", "", "humour", "testimonial", "customer", "ugc", "no", "", "no", "english_first", "Repost @Sanober Yasir - two women, sleepwear set."),
"Dbd7nn6sGLI": ("Micellar Cleansing Oil-in-Water", "genz", "ugc_repost", "cleansing_purity", "", "humour", "testimonial", "customer", "ugc", "no", "", "no", "english_first", "Repost @B O B B Y Official - barber's chair. The only male-cast post across the three multinationals."),
"DbdtCW4oqY8": ("Micellar Cleansing Oil-in-Water", "genz", "ugc_repost", "cleansing_purity", "", "aspiration", "testimonial", "customer", "ugc", "no", "", "yes", "english_first", "Repost @SARA ALI - makeup-removal demo."),
"DbkmvQ7iJ5w": ("Micellar Cleansing Oil-in-Water", "unclear", "brand_building", "other", "", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "'IF YOU NEED....' over pink water. Teaser."),
# ---------------- Nivea Pakistan (all 2022) ----------------
"CeGrP36tqcA": ("NIVEA Soft", "unclear", "product_feature", "hydration_moisture", "", "reassurance", "none", "none_people", "studio", "no", "", "no", "english_first", "'Experience the richness of NIVEA Soft'."),
"CeL0yiYt4m-": ("NIVEA Soft", "geny", "product_feature", "hydration_moisture", "", "aspiration", "none", "customer", "studio", "no", "", "no", "english_first", "'Instant Softness!' - model holds the tin."),
"CeR8uOmovcE": ("Body lotion range", "genz", "brand_building", "hydration_moisture", "", "humour", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "Mock 'Suggestions for you' Instagram panel: Shea smooth / Cocoa Butter / Rich Nourishing, each with a Follow button."),
"CeWIBd8t-aC": ("Natural Fairness body lotion", "unclear", "product_feature", "brightening_tone", "", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "'Bring out your natural radiance'."),
"Ced5gJfoGzy": ("Powder Touch deodorant", "unclear", "product_feature", "other", "", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "Deodorant, not skincare. Roll-on in pink powder."),
"CegbKSPt_dg": ("Fresh Natural deodorant", "genz", "product_feature", "other", "", "reassurance", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "Mock search bar: 'How to stay fresh during summers'. Deodorant, not skincare."),
"CghK6Mbte8K": ("NIVEA Soft", "geny", "brand_building", "hydration_moisture", "", "reassurance", "none", "customer", "studio", "no", "", "no", "english_first", "'Show your skin some love' inside a retro desktop-window frame."),
"Cgjvx9Zt-9l": ("Aloe & Hydration body lotion", "unclear", "product_feature", "hydration_moisture", "", "reassurance", "claim_statistics", "none_people", "studio", "no", "", "no", "english_first", "'Aloe & Hydration 48h'."),
"CgmUd8kNFFG": ("Powder Touch deodorant", "unclear", "product_feature", "other", "", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "Spray + roll-on on a plinth. Deodorant."),
"Cgo5TxItfiE": ("NIVEA Soft", "unclear", "product_feature", "ingredient_science", "hydration_moisture", "reassurance", "none", "none_people", "studio", "no", "", "yes", "english_first", "'Jojoba Oil / Vitamin E / Fast absorbing / Instant soft skin'."),
"CgreJurt0y7": ("NIVEA Soft", "geny", "product_feature", "hydration_moisture", "", "reassurance", "none", "customer", "studio", "no", "", "no", "english_first", "Model in a towel turban holding the tin."),
"CguC8SFNwrf": ("Body lotion range", "unclear", "product_feature", "natural_ingredients", "hydration_moisture", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "Berries-and-lotion flat lay."),
# ---------------- Saeed Ghani ----------------
"Db5qUcXNld9": ("Detox & Cleanse Activated Charcoal Face Wash", "genz", "product_feature", "acne_oil_control", "", "reassurance", "testimonial", "customer", "ugc", "no", "Refreshed & Clear Skin with every wash", "yes", "english_first", "Creator collage: 'When your skin gets the detox it needs, it shows'."),
"Db_FijUNTfF": ("Detox & Cleanse Activated Charcoal Face Wash", "unclear", "product_feature", "acne_oil_control", "", "aspiration", "none", "none_people", "studio", "no", "Refreshed & Clear Skin with every wash", "no", "english_first", "'SKIN DETOX POWER' with www.saeedghani.pk burnt into the artwork."),
"DcBc0xvje9K": ("Detox & Cleanse Activated Charcoal Face Wash", "geny", "product_feature", "acne_oil_control", "", "reassurance", "none", "customer", "stock", "no", "", "yes", "english_first", "'IF YOUR SKIN FEELS OILY & DIRTY THIS SUMMER! Try This' over a man's face against sky. Male casting."),
"DcGzXLDhbo3": ("Detox & Cleanse Activated Charcoal Face Wash", "geny", "brand_building", "cleansing_purity", "", "aspiration", "none", "none_people", "studio", "no", "", "no", "english_first", "'Deep Clean Reset' - tube on a stool in a gym."),
"DcLrpncNgbF": ("Detox & Cleanse Activated Charcoal Face Wash", "unclear", "product_feature", "acne_oil_control", "", "reassurance", "testimonial", "none_people", "template_graphic", "no", "", "yes", "english_first", "Five-star review card, verbatim: 'I just used it from three days my skin has become just super cool so i m ordering this by my heart'."),
"DcOSE-EtN_V": ("Detox & Cleanse Activated Charcoal Face Wash", "unclear", "product_feature", "acne_oil_control", "", "reassurance", "none", "none_people", "studio", "no", "Refreshed & Clear Skin with every wash", "yes", "english_first", "'Controls oil / Detoxifies skin / Unclogs pores'."),
"DcQyLM-tjJs": ("Detox & Cleanse Activated Charcoal Face Wash", "unclear", "product_feature", "acne_oil_control", "", "reassurance", "none", "none_people", "studio", "no", "", "yes", "english_first", "'SKIN DETOX - Controls oil | Detoxifies skin | Unclogs pores'."),
"DcTcnxkN4Yb": ("Detox & Cleanse Activated Charcoal Face Wash", "unclear", "product_feature", "ingredient_science", "acne_oil_control", "reassurance", "none", "none_people", "studio", "no", "", "yes", "english_first", "'HERO INGREDIENTS - Activated Charcoal / Menthol'."),
"DcWRgxojazA": ("Detox & Cleanse Activated Charcoal Face Wash", "geny", "product_feature", "acne_oil_control", "", "reassurance", "none", "customer", "stock", "no", "", "yes", "english_first", "Bearded man, half-desaturated, skin-zoom circle: 'If your skin feels oily & dirty this summer'."),
"DcbS55dtmUH": ("Detox & Cleanse Activated Charcoal Face Wash", "geny", "product_feature", "ingredient_science", "acne_oil_control", "reassurance", "none", "customer", "studio", "no", "Refreshed & Clear Skin with every wash", "yes", "english_first", "'DEEP CLEANSING POWER - Activated Charcoal / Menthol'."),
"DcdrEx5t4zg": ("Detox & Cleanse Activated Charcoal Face Wash", "unclear", "product_feature", "acne_oil_control", "", "reassurance", "none", "none_people", "studio", "no", "Refreshed & Clear Skin with every wash", "yes", "english_first", "'Reveal your skin's glow ... Clears dirt & dullness, Instantly Cools & Refreshes'."),
"DcggNGzNX7C": ("Detox & Cleanse Activated Charcoal Face Wash", "genz", "product_feature", "acne_oil_control", "", "humour", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "Chat-bubble mock: 'Hey, I got you something!' / 'OMG! My skin needed this!!!'"),
# ---------------- Dermaceutical Pakistan ----------------
"DGrpObpIGOj": ("Niacinamide 10% serum", "geny", "product_feature", "ingredient_science", "", "reassurance", "testimonial", "customer", "ugc", "no", "", "yes", "english_first", "'7 Day Challenge 10% Niacinamide' - older pinned post (Mar 2025)."),
"DJTYcqZsR6j": ("unclear", "geny", "product_feature", "other", "", "reassurance", "none", "customer", "ugc", "no", "", "no", "other", "Older pinned post (May 2025). No copy legible at thumbnail."),
"DKFFLfcTdoe": ("Dermaceutical Facial Kit", "genz", "ugc_repost", "other", "", "humour", "testimonial", "customer", "ugc", "no", "", "no", "english_first", "Repost @Aymen Zahra: 'DERMACEUTICAL FACIAL KIT' (May 2025)."),
"DcGjsBvMasf": ("Tinted Sunscreen SPF50", "geny", "product_feature", "sun_protection", "", "reassurance", "none", "customer", "template_graphic", "no", "One Step. Many Benefits", "yes", "english_first", "'TINTED SUNSCREEN With Light Foundation ... ONE STEP. MANY BENEFITS' - dense spec sheet."),
"DcVRtShMrwS": ("Serum + sunscreen range", "unclear", "tactical_promo", "price_value", "", "pride", "none", "none_people", "template_graphic", "no", "", "no", "equal", "'RABI-UL-AWWAL HOLY MONTH SALE - 10% OFF on all serums & tinted mineral sunscreen'."),
"DcVzfRZsdUQ": ("brand", "genz", "ugc_repost", "brand_corporate", "", "pride", "testimonial", "customer", "ugc", "no", "", "yes", "english_first", "Repost @Paro: 'THE ONLY PAKISTANI SKINCARE BRAND I TRUST!!'"),
"Dcby3nso8QX": ("Tinted Sunscreen SPF50", "genz", "ugc_repost", "sun_protection", "", "aspiration", "testimonial", "none_people", "ugc", "no", "", "no", "english_first", "Repost: 'POV: You finally found your everyday SPF'."),
"Dcc5xQatSBg": ("Alpha Arbutin serum", "genz", "ugc_repost", "brightening_tone", "ingredient_science", "reassurance", "expert_authority", "customer", "ugc", "no", "", "yes", "english_first", "Repost @Doc.Rapunzel: notes-app framing, 'Brightening Serum - Alpha Arbutin'."),
"DcesPj0smw2": ("Tinted Sunscreen SPF50", "unclear", "product_feature", "sun_protection", "", "aspiration", "none", "none_people", "template_graphic", "no", "", "no", "english_first", "'LAUNCHING SOON - TINTED SUNSCREEN SPF 50 PA++++'."),
"DcgLVf7NYV7": ("Serum range", "geny", "ugc_repost", "acne_oil_control", "", "reassurance", "testimonial", "customer", "ugc", "no", "", "yes", "english_first", "Repost @Doc.Rapunzel: 'Facing acne Problems' with a before/after face grid. The only visible before/after in the whole scan."),
"DcgQz5qTwfz": ("Serum range", "geny", "ugc_repost", "other", "", "humour", "testimonial", "customer", "ugc", "no", "", "no", "english_first", "Repost @Hira Zahid Butt: 'WHEN HE STEALS YOUR BEAUTY MUST-HAVES'."),
"DcgSevHiNce": ("Tinted Sunscreen SPF50", "geny", "ugc_repost", "sun_protection", "", "reassurance", "testimonial", "customer", "ugc", "no", "", "yes", "english_first", "Repost @Rija Rafique: 'Real Results of Influencer-recommended Tinted SPF'."),
# ---------------- Conatural ----------------
"DcLXruZgaun": ("NO B.S. (brand platform)", "unclear", "brand_building", "brand_corporate", "ingredient_science", "duty", "none", "none_people", "template_graphic", "no", "NO B.S.", "yes", "english_first", "OFFICIAL STATEMENT: 'We are aware that our recent posts have caused a lot of outrage. And honestly, we are glad they did. Because false, and fabricated claims have no space in our industry... No misleading labels. No ingredient lies. No promises that sound too good to be true.'"),
"DcNfqpfIiTY": ("NO B.S. (brand platform)", "genz", "brand_building", "other", "", "humour", "none", "customer", "studio", "no", "NO B.S.", "no", "english_first", "Woman seated under a wall-painted 'no B.S.'"),
"DcOkR7tof1X": ("NO B.S. (brand platform)", "unclear", "brand_building", "other", "", "humour", "none", "none_people", "template_graphic", "no", "NO B.S.", "no", "english_first", "Teaser: 'CONATURAL no B.S. (bad ____)' with the noun blanked out."),
"DcQbKPZglAW": ("Super Activs", "unclear", "brand_building", "ingredient_science", "", "duty", "none", "none_people", "template_graphic", "no", "NO B.S.", "yes", "english_first", "'We put ACTIVS in. Not PROMISES.'"),
"DcQbVltAo4u": ("Super Activs Niacinamide 10% & Zinc 1%", "genz", "product_feature", "ingredient_science", "", "humour", "none", "customer", "studio", "no", "NO B.S.", "yes", "english_first", "'no B.S. (bad skin)' with the serum held to camera."),
"DcQkMcPA1o9": ("Super Activs Niacinamide 10% & Zinc 1%", "genz", "brand_building", "ingredient_science", "", "humour", "none", "none_people", "studio", "no", "NO B.S.", "no", "english_first", "'Activs in. Drama out.'"),
"DcWsTb-ItrU": ("NO B.S. (brand platform)", "unclear", "brand_building", "ingredient_science", "", "duty", "none", "none_people", "template_graphic", "no", "NO B.S.", "yes", "english_first", "'NOT A MIRACLE. Just good ingredients doing their job.'"),
"DcWsZPqoDtC": ("Tea Tree & Neem Face Wash", "unclear", "product_feature", "natural_ingredients", "", "humour", "none", "none_people", "studio", "no", "NO B.S.", "no", "english_first", "'Start clean. Stay clean. No B.S. in between.'"),
"DcYi48Co84t": ("Tea Tree & Neem Face Wash", "genz", "product_feature", "natural_ingredients", "", "humour", "none", "customer", "studio", "no", "NO B.S.", "no", "english_first", "'no B.S. (bad skin)' with the face wash."),
"Dca4sGMg6kf": ("NO B.S. (brand platform)", "unclear", "brand_building", "sensitivity_safety", "", "duty", "none", "none_people", "template_graphic", "no", "NO B.S.", "yes", "english_first", "'No Sulphates No Parabens NO B.S.'"),
"Dca4v08gReZ": ("Rosemary & Onion Shampoo (haircare)", "geny", "product_feature", "natural_ingredients", "", "humour", "none", "customer", "studio", "no", "NO B.S.", "no", "english_first", "'no B.S. (bad scalp)'. Haircare, not skincare - the platform is stretched across the whole portfolio."),
"DcePcXUoyRi": ("Shampoo range (haircare)", "unclear", "product_feature", "natural_ingredients", "", "humour", "none", "none_people", "studio", "no", "NO B.S.", "no", "english_first", "'No B.S. From roots to ends.' Haircare."),
}

SKINCARE_CLAIMS = ["brightening_tone", "hydration_moisture", "sun_protection", "acne_oil_control",
                   "anti_ageing", "natural_ingredients", "ingredient_science", "cleansing_purity",
                   "sensitivity_safety", "price_value", "brand_corporate", "other"]


def main():
    posts = list(csv.DictReader((WS / "posts.csv").open(encoding="utf-8")))
    handles = {r["brand"]: r for r in csv.DictReader((WS / "handles.csv").open(encoding="utf-8-sig"))}

    wb = openpyxl.load_workbook(WS / "capture.xlsx")
    ads = wb["Ads"]
    COLS = [c.value for c in ads[1]]
    idx = {c: i for i, c in enumerate(COLS)}

    # rewrite the bound claim vocabulary for skincare (SKILL.md requires this per category)
    codes = wb["Codes"]
    claim_col = next(c.column_letter for c in codes[1] if c.value == "claim")
    for j in range(2, 40):
        codes[f"{claim_col}{j}"] = None
    for j, v in enumerate(SKINCARE_CLAIMS, start=2):
        codes[f"{claim_col}{j}"] = v

    n = 0
    missing = []
    for p in posts:
        sc = p["shortcode"]
        if sc not in C:
            missing.append(sc)
            continue
        (focus, gen, ctype, cl1, cl2, reg, proof, frame, prod, eco, plat, obj, lang, note) = C[sc]
        m = re.search(r" on (\w+ \d+, \d{4})", p["alt_text"])
        d = datetime.strptime(m.group(1), "%B %d, %Y").date().isoformat() if m else ""
        row = [""] * len(COLS)
        row[idx["post_id"]] = sc
        row[idx["brand"]] = p["brand"]
        row[idx["ring"]] = p["ring"]
        row[idx["handle"]] = p["handle"]
        row[idx["date"]] = d
        row[idx["format"]] = "video" if p["kind"] == "reel" else "static"
        row[idx["product_focus"]] = focus
        row[idx["audience_generation"]] = gen
        row[idx["content_type"]] = ctype
        row[idx["claim_primary"]] = cl1
        row[idx["claim_secondary"]] = cl2
        row[idx["register"]] = reg
        row[idx["proof_device"]] = proof
        row[idx["who_is_in_frame"]] = frame
        row[idx["production_value"]] = prod
        row[idx["ecosystem_push"]] = eco
        row[idx["campaign_platform"]] = plat
        row[idx["addresses_objection"]] = obj
        row[idx["language_lead"]] = lang
        row[idx["reposted_from"]] = p["reposted_from"]
        row[idx["source_url"]] = p["source_url"]
        row[idx["image_file"]] = p["image_file"]
        row[idx["coder"]] = CODER
        row[idx["notes"]] = note
        ads.append(row)
        n += 1

    br = wb["Brands"]
    prof = {r["brand"]: r for r in csv.DictReader((WS / "profiles.csv").open(encoding="utf-8"))}
    for r in br.iter_rows(min_row=2):
        b = r[0].value
        if b in handles:
            r[2].value = handles[b]["handle"]
        if b in prof:
            r[3].value = prof[b]["followers"]

    rd = wb["README"]
    for row in rd.iter_rows(min_row=1, max_row=rd.max_row):
        if row[0].value == "Core category objection":
            row[1].value = ('"Does it actually do anything to my skin, or is it just a '
                            'brightening promise printed on a tube?"')
    rd.append(["Claim vocabulary", "REWRITTEN for facial skincare: " + " / ".join(SKINCARE_CLAIMS)])
    rd.append(["Absences", "Cetaphil Pakistan and Olay Pakistan have no Pakistan-specific Instagram account. Recorded, not dropped."])

    wb.save(WS / "capture.xlsx")
    sys.stdout.reconfigure(encoding="utf-8")
    print(f"coded {n} rows; unmatched shortcodes: {missing}")


if __name__ == "__main__":
    main()
