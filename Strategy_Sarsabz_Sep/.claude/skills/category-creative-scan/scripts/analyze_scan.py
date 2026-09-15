"""Turn the coded capture sheet into a findings report and CSVs.

    python analyze_scan.py <workspace>/scan/capture.xlsx --out <workspace>/scan

Produces scan_findings.md plus: product_focus.csv, generation.csv, content_balance.csv,
cadence.csv, territory_map.csv, proof.csv, frame.csv.

The tables are inputs, not the finding. The finding is one sentence naming the vacant
territory and one paragraph proving it - you write that.
"""

from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import sys

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def load(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Ads"]
    hdr = [c.value for c in ws[1]]
    rows = []
    dropped = 0
    for r in ws.iter_rows(min_row=2):
        d = dict(zip(hdr, [c.value for c in r]))
        if not d.get("brand"):
            continue
        if not d.get("source_url"):
            dropped += 1
            continue
        rows.append(d)
    if dropped:
        print(f"  dropped {dropped} row(s) with no source_url")
    return rows


def market_share(path: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "MarketShare" not in wb.sheetnames:
        return {}
    ws = wb["MarketShare"]
    hdr = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2):
        d = dict(zip(hdr, [c.value for c in r]))
        if d.get("brand") and d.get("market_share_pct") not in (None, ""):
            out[d["brand"]] = d
    return out


def table(add, title, header, rows, note: str | None = None) -> None:
    add(f"## {title}\n")
    if note:
        add(note + "\n")
    if not rows:
        add("_No data coded for this dimension yet._\n")
        return
    add("| " + " | ".join(header) + " |")
    add("|" + "|".join("---" for _ in header) + "|")
    for r in rows:
        add("| " + " | ".join(str(x) for x in r) + " |")
    add("")


def pct(n: int, d: int) -> str:
    return f"{round(100*n/d)}%" if d else "-"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("capture", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    out = args.out
    out.mkdir(parents=True, exist_ok=True)

    rows = load(args.capture)
    if not rows:
        raise SystemExit("No coded rows with a source_url.")
    shares = market_share(args.capture)

    brands = sorted({r["brand"] for r in rows})
    by = defaultdict(list)
    for r in rows:
        by[r["brand"]].append(r)

    L: list[str] = []
    add = L.append
    add("# Category creative scan - findings\n")
    add(f"Source: `{args.capture.name}`. Generated {date.today().isoformat()}.\n")

    # ---- 1. sample frame ----
    add("## 1. What we looked at\n")
    add(f"**{len(rows)} posts** coded across **{len(brands)} brands**.\n")
    counts = Counter(r["brand"] for r in rows)
    equal = len(set(counts.values())) == 1 and len(counts) > 1
    table(add, "Sample frame", ["Brand", "Ring", "Posts"],
          [(b, (by[b][0].get("ring") or "-"), counts[b]) for b in brands])
    if equal:
        add("> **Every brand contributed the same number of posts.** That is a collection\n"
            "> artefact of the capped sample, not a fact about the category. **Share of voice\n"
            "> cannot be measured from this data and must not be presented.**\n")
    unclear = sum(1 for r in rows if "unclear" in
                  {r.get("audience_generation"), r.get("production_value")})
    if unclear:
        add(f"{unclear} post(s) carry an `unclear` code on at least one dimension - "
            f"coded honestly rather than guessed.\n")

    # ---- 2. product focus / convergence ----
    focus = defaultdict(Counter)
    for r in rows:
        f = (r.get("product_focus") or "none").strip()
        focus[r["brand"]][f] += 1
    rows_pf = []
    for b in brands:
        top = focus[b].most_common(3)
        rows_pf.append((b, "; ".join(f"{k} ({v})" for k, v in top) or "-"))
    table(add, "2. What each brand is pushing", ["Brand", "Top product focus"], rows_pf,
          "Convergence is the thing to look for: several brands funnelling every message "
          "into an owned app or ecosystem while their claims look different on a message map.")

    eco = [(b, sum(1 for r in by[b] if str(r.get("ecosystem_push")).lower() == "yes"), len(by[b]))
           for b in brands]
    eco = [(b, n, t, pct(n, t)) for b, n, t in eco if n]
    if eco:
        table(add, "2a. Ecosystem convergence", ["Brand", "Ads driving to an owned app", "Of", "Share"],
              sorted(eco, key=lambda x: -x[1]),
              "How much of the category's output exists to drive app installs rather than "
              "to sell the category's actual product.")

    # ---- 3. audience generation ----
    gens = ["genz", "geny", "genx", "boomer_plus", "mixed", "unclear"]
    present = [g for g in gens if any(r.get("audience_generation") == g for r in rows)]
    table(add, "3. Who the work is cast for", ["Brand"] + present,
          [[b] + [sum(1 for r in by[b] if r.get("audience_generation") == g) or "-"
                  for g in present] for b in brands],
          "**Inference from casting, styling, setting and language - not a fact about who "
          "buys.** Label it as inference wherever it is presented. A generation the whole "
          "category has stopped addressing is a gap worth naming.")

    # ---- 4. content balance ----
    types = ["brand_building", "tactical_promo", "product_feature", "corporate_pr",
             "recruitment", "calendar_topical", "csr", "ugc_repost"]
    present_t = [t for t in types if any(r.get("content_type") == t for r in rows)]
    table(add, "4. Brand-building vs tactical vs PR", ["Brand"] + present_t + ["Building %"],
          [[b] + [sum(1 for r in by[b] if r.get("content_type") == t) or "-" for t in present_t]
           + [pct(sum(1 for r in by[b] if r.get("content_type") == "brand_building"), len(by[b]))]
           for b in brands],
          "The honesty test. A brand can look busy while publishing nothing that builds it.")

    cal = [(b, sum(1 for r in by[b] if r.get("content_type") == "calendar_topical"),
            len(by[b])) for b in brands]
    cal = [(b, k, t, pct(k, t)) for b, k, t in cal if k]
    if cal:
        table(add, "4b. Output that sells nothing", ["Brand", "Calendar posts", "Of", "Share"],
              sorted(cal, key=lambda x: -x[1]),
              "Independence days, religious greetings, weather advisories. A brand whose "
              "feed is half calendar content is not running a campaign, it is filling one.")

    reposts = [(b, sum(1 for r in by[b] if r.get("reposted_from")), len(by[b])) for b in brands]
    reposts = [(b, n, t, pct(n, t)) for b, n, t in reposts if n]
    if reposts:
        table(add, "4a. Feeds padded with creator content", ["Brand", "Reposts", "Of", "Share"],
              sorted(reposts, key=lambda x: -x[1]),
              "Reposted creator content is not the brand's own creative.")

    # ---- 5. cadence ----
    cad = []
    pinned_warned = False
    for b in brands:
        ds = []
        # Instagram serves up to 3 PINNED posts at the top of a grid, and a pinned post
        # can be months older than everything else. Taking min/max across the grid then
        # reports a brand that published 5 posts in 6 days as 0.6 posts/week. Exclude
        # flagged pins from the span; they still count as published output elsewhere.
        pool = [r for r in by[b] if str(r.get("pinned_guess", "")).lower() != "yes"] or by[b]
        if len(pool) < len(by[b]):
            pinned_warned = True
        for r in pool:
            v = r.get("date")
            if isinstance(v, datetime):
                ds.append(v.date())
            elif isinstance(v, date):
                ds.append(v)
            elif isinstance(v, str) and v.strip():
                try:
                    ds.append(datetime.strptime(v.strip()[:10], "%Y-%m-%d").date())
                except ValueError:
                    pass
        if len(ds) < 2:
            cad.append((b, len(pool), "-", "-", "insufficient dates"))
            continue
        raw_span = (max(ds) - min(ds)).days
        since = (date.today() - max(ds)).days
        # A span shorter than a week turns a burst into a nonsense annualised rate -
        # nine posts in one day is not 63 posts/week. Floor the span and say so.
        span = max(raw_span, 7)
        per_week = len(ds) / span * 7
        if since > 60:
            # Recency beats rate. A brand that published steadily and then stopped a year
            # ago is dormant, however busy the surviving posts make it look.
            band = "DORMANT"
        elif per_week >= 4:
            band = "active"
        elif per_week >= 1.5:
            band = "moderate"
        else:
            band = "slow"
        if raw_span < 7:
            band += " (burst)"
        cad.append((b, len(ds), f"{per_week:.1f}", since, band))
    table(add, "5. Cadence", ["Brand", "Posts dated", "Posts/week", "Days since last", "Band"],
          sorted(cad, key=lambda x: (x[2] == "-", -float(x[2]) if x[2] != "-" else 0)),
          "Computed from post dates, never eyeballed. With a capped sample this is a rate "
          "over the recent window, not a long-run average."
          + (" Pinned posts are excluded from the span - a pin can be months old and would "
             "otherwise make an active brand read as dormant." if pinned_warned else "")
          + " `DORMANT` means the last post is over 60 days old, whatever the rate says. "
            "`(burst)` means every dated post fell inside one week, so the rate is a "
            "spike, not a habit.")

    # ---- 6. territory map ----
    claims = sorted({r.get("claim_primary") or "-" for r in rows})
    regs = sorted({r.get("register") or "-" for r in rows})
    grid = Counter((r.get("claim_primary") or "-", r.get("register") or "-") for r in rows)
    empty = sum(1 for c in claims for g in regs if not grid[(c, g)])
    table(add, "6. Message territory map", ["Claim \\ Register"] + regs + ["Total"],
          [[c] + [grid[(c, g)] or "-" for g in regs] + [sum(grid[(c, g)] for g in regs)]
           for c in claims],
          f"**{empty} of {len(claims)*len(regs)} claim/register combinations are unoccupied.** "
          "Some are empty because they do not work; the defensible territory is usually among "
          "the rest.")

    # ---- 7. proof and objection ----
    pf = Counter(r.get("proof_device") or "none" for r in rows)
    table(add, "7. Proof devices", ["Device", "Ads", "Share"],
          [(k, v, pct(v, len(rows))) for k, v in pf.most_common()],
          "Where most categories are hollow. Distinguish `claim_statistics` (a performance "
          "number) from `ratings_awards` (someone else's badge) - the first is far harder to copy.")
    obj = sum(1 for r in rows if str(r.get("addresses_objection")).lower() == "yes")
    add(f"**Ads addressing the category's core objection head-on: {obj} of {len(rows)} "
        f"({pct(obj, len(rows))}).**\n")

    frame = Counter(r.get("who_is_in_frame") or "-" for r in rows)
    cust = frame.get("customer", 0)
    exe = frame.get("executive", 0)
    table(add, "8. Who is in frame", ["Who", "Ads", "Share"],
          [(k, v, pct(v, len(rows))) for k, v in frame.most_common()],
          f"Customer-to-executive ratio across the category: **{cust}:{exe}**. A grid "
          "dominated by executives is a diagnosis, not a style choice.")

    # ---- script, measured on the RAW caption text, never on our own coding.
    # `language_lead` has a default value, so counting it measures the coder, not the
    # category. This reads the captions the collector saved and counts actual scripts.
    posts_csv = out / "posts.csv"
    if posts_csv.exists():
        import re as _re
        NONLATIN = {
            "Arabic/Urdu": _re.compile(r"[؀-ۿ]"),
            "Devanagari": _re.compile(r"[ऀ-ॿ]"),
            "CJK": _re.compile(r"[一-鿿]"),
            "Cyrillic": _re.compile(r"[Ѐ-ӿ]"),
        }
        with posts_csv.open(encoding="utf-8", newline="") as fh:
            caps = [(r.get("brand", ""), r.get("caption") or "")
                    for r in csv.DictReader(fh)]
        captioned = [(b, c) for b, c in caps if c.strip()]
        script_hits = Counter()
        by_brand_script = defaultdict(int)
        for b, c in captioned:
            for name, rx in NONLATIN.items():
                if rx.search(c):
                    script_hits[name] += 1
                    by_brand_script[b] += 1
                    break
            else:
                script_hits["Latin only"] += 1
        if captioned:
            table(add, "9a. Script actually used in captions",
                  ["Script", "Posts", "Share of captioned"],
                  [(k, v, pct(v, len(captioned))) for k, v in script_hits.most_common()],
                  f"Measured on the raw caption text of {len(captioned)} captioned posts, "
                  "independent of any coding - a `language_lead` default would otherwise "
                  "measure the coder rather than the category. In a market whose national "
                  "language is not Latin-script, a category advertising almost entirely in "
                  "Latin script has chosen an audience, whether or not it meant to.")
            if by_brand_script:
                table(add, "9b. Brands using a non-Latin script at all",
                      ["Brand", "Posts"],
                      sorted(by_brand_script.items(), key=lambda kv: -kv[1]))

    prod = Counter(r.get("production_value") or "-" for r in rows)
    table(add, "9. Production value", ["Type", "Ads", "Share"],
          [(k, v, pct(v, len(rows))) for k, v in prod.most_common()],
          "A brand running template graphics for twelve straight posts is not running a "
          "campaign, whatever its captions claim.")

    inv = [(b,
            sum(1 for r in by[b] if r.get("production_value") == "studio"),
            sum(1 for r in by[b] if r.get("who_is_in_frame") == "customer"),
            sum(1 for r in by[b] if r.get("who_is_in_frame") in ("executive", "staff")))
           for b in brands]
    table(add, "9c. Investment, and who appears in the work",
          ["Brand", "Studio shoots", "Customer in frame", "Executive / staff"],
          sorted(inv, key=lambda x: -x[1]),
          "Commissioned photography is the clearest signal of what a brand is actually "
          "spending. Read it against who is on screen: a brand that shoots often and casts "
          "its own executives is investing in itself, not in its customer.")

    plats = defaultdict(Counter)
    for r in rows:
        p = (r.get("campaign_platform") or "").strip()
        if p:
            plats[r["brand"]][p] += 1
    if plats:
        table(add, "10. Recurring campaign platforms", ["Brand", "Line / hashtag", "Ads"],
              [(b, k, v) for b in brands for k, v in plats[b].most_common(2)],
              "A recurring line is a platform; a one-off is not. Brands with one are much "
              "harder to displace than their follower counts suggest.")

    if shares:
        table(add, "11. Presence vs market share", ["Brand", "Posts", "Market share", "Source"],
              [(b, counts[b], shares[b].get("market_share_pct"), shares[b].get("source") or "-")
               for b in brands if b in shares],
              "Compare *presence* with position. Not spend - this sample cannot measure spend.")
    else:
        add("## 11. Presence vs market share\n")
        add("> `MarketShare` sheet is empty. Fill it with cited regulator or annual-report\n"
            "> figures. Even without share of voice, position-vs-presence is worth showing.\n")

    coders = Counter(r.get("coder") or "-" for r in rows)
    if len(coders) > 1:
        add("## Coder drift check\n")
        add("| Coder | Ads | brand_building % | proof=none % |")
        add("|---|---:|---:|---:|")
        for c, n in coders.most_common():
            rs = [r for r in rows if (r.get("coder") or "-") == c]
            add(f"| {c} | {n} | "
                f"{pct(sum(1 for r in rs if r.get('content_type')=='brand_building'), len(rs))} | "
                f"{pct(sum(1 for r in rs if r.get('proof_device')=='none'), len(rs))} |")
        add("\nSharp divergence on one dimension means recode it rather than shipping mixed data.\n")

    add("## Next\n")
    add("These tables are inputs. Write one sentence naming the vacant territory and one\n"
        "paragraph proving it from this evidence, then state what the client can credibly\n"
        "own - anchored in an asset competitors structurally cannot copy. If the analysis\n"
        "could be deleted without weakening the campaign concept, it has failed.\n")

    (out / "scan_findings.md").write_text("\n".join(L), encoding="utf-8")

    def dump(name, header, data):
        with (out / name).open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(header)
            w.writerows(data)

    dump("product_focus.csv", ["brand", "product_focus", "ads"],
         [(b, k, v) for b in brands for k, v in focus[b].most_common()])
    dump("generation.csv", ["brand"] + present,
         [[b] + [sum(1 for r in by[b] if r.get("audience_generation") == g) for g in present]
          for b in brands])
    dump("content_balance.csv", ["brand"] + present_t,
         [[b] + [sum(1 for r in by[b] if r.get("content_type") == t) for t in present_t]
          for b in brands])
    dump("cadence.csv", ["brand", "posts_dated", "posts_per_week", "days_since_last", "band"], cad)
    dump("territory_map.csv", ["claim"] + regs,
         [[c] + [grid[(c, g)] for g in regs] for c in claims])
    dump("proof.csv", ["proof_device", "ads"], pf.most_common())
    dump("frame.csv", ["who_is_in_frame", "ads"], frame.most_common())

    print(f"Wrote {out/'scan_findings.md'}")
    print("       + product_focus, generation, content_balance, cadence, territory_map, proof, frame (CSV)")
    print(f"Analysed {len(rows)} posts across {len(brands)} brands.")


if __name__ == "__main__":
    main()
