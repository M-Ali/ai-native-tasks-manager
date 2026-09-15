"""Apply codes from a plain CSV into capture.xlsx, validating against the vocabularies.

    python apply_codes.py <workspace>/scan/codes.csv --workspace <workspace>/scan

`codes.csv` needs a `post_id` column plus any columns you want to set. Everything else
is ignored, so you can code in whatever tool you like — a spreadsheet, an editor, or a
script — and land it here safely:

    post_id,product_focus,audience_generation,content_type,who_is_in_frame
    DaKLEGSh1X5,brand,geny,brand_building,customer
    DcaZMaQiguj,Jubilee Active,genz,tactical_promo,customer

Why this exists: coding 150+ posts inline in a script hits shell argument limits on
Windows, and hand-editing the workbook loses provenance. A CSV is the smallest thing
that works.

Validates every value against the controlled vocabulary BEFORE writing anything. A typo
fails here with a list of what was wrong, rather than silently poisoning the analysis —
a bad code is worse than a blank one because it looks like data.
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, str(Path(__file__).parent))
from make_handles import BIND, VOCAB  # noqa: E402  single-sourced vocabulary

FREE_TEXT = {"product_focus", "campaign_platform", "notes", "coder"}


def vocab_from(wb) -> dict[str, list[str]]:
    """The workbook's Codes sheet is the vocabulary of record for THIS scan.

    make_handles writes it there, including a claim list replaced with --claims for a
    non-financial category. Validating against the script's defaults instead would
    reject the very codes the dropdowns offer.
    """
    if "Codes" not in wb.sheetnames:
        return VOCAB
    ws = wb["Codes"]
    found = {}
    for col in ws.iter_cols(values_only=True):
        if not col or not col[0]:
            continue
        vals = [str(v).strip() for v in col[1:] if v is not None and str(v).strip()]
        if vals:
            found[str(col[0]).strip()] = vals
    missing = [k for k in VOCAB if k not in found]
    if missing:
        print(f"  Codes sheet has no column for {', '.join(missing)} - using script defaults")
    return {**VOCAB, **found}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("codes", type=Path)
    ap.add_argument("--workspace", type=Path, required=True)
    ap.add_argument("--capture", type=Path, default=None)
    args = ap.parse_args()

    cap_path = args.capture or args.workspace / "capture.xlsx"
    if not cap_path.exists():
        sys.exit(f"{cap_path} not found.")

    with args.codes.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"{args.codes} has no rows.")
    if "post_id" not in rows[0]:
        sys.exit(f"{args.codes} needs a post_id column.")

    wb = openpyxl.load_workbook(cap_path)
    vocab = vocab_from(wb)
    ws = wb["Ads"]
    hdr = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(hdr)}

    fields = [f for f in rows[0] if f != "post_id" and f in col]
    unknown = [f for f in rows[0] if f != "post_id" and f not in col]
    if unknown:
        print(f"  ignoring column(s) not in the sheet: {', '.join(unknown)}")
    if not fields:
        sys.exit("No codeable columns in the CSV match the sheet.")

    # validate everything before touching the workbook
    errors = []
    for r in rows:
        for f in fields:
            v = (r.get(f) or "").strip()
            if not v or f in FREE_TEXT:
                continue
            vocab_key = BIND.get(f)
            if vocab_key and v not in vocab[vocab_key]:
                errors.append(f"{r['post_id']}: {f}={v!r} not in "
                              f"[{', '.join(vocab[vocab_key])}]")
    if errors:
        for e in errors[:25]:
            print(f"  {e}")
        if len(errors) > 25:
            print(f"  ... and {len(errors)-25} more")
        sys.exit(f"\n{len(errors)} invalid code(s). Nothing written.")

    by_id = {str(r["post_id"]).strip(): r for r in rows}
    updated = cells = 0
    for excel_row in range(2, ws.max_row + 1):
        pid = ws.cell(row=excel_row, column=col["post_id"]).value
        if not pid:
            continue
        src = by_id.get(str(pid).strip())
        if not src:
            continue
        touched = False
        for f in fields:
            v = (src.get(f) or "").strip()
            if v:
                ws.cell(row=excel_row, column=col[f], value=v)
                cells += 1
                touched = True
        if touched:
            updated += 1

    wb.save(cap_path)
    unmatched = len(by_id) - updated
    print(f"{cap_path}")
    print(f"  {updated} row(s) updated, {cells} cell(s) written across {len(fields)} field(s)")
    if unmatched > 0:
        print(f"  {unmatched} post_id(s) in the CSV matched no row in the sheet - "
              f"check they came from this collection")


if __name__ == "__main__":
    main()
