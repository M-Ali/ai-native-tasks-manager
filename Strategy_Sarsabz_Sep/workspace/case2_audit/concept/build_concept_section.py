"""Campaign concept section for the Sarsabz Case 2 deck, rendered from concept.yaml.

    uv run --with python-pptx --with pyyaml python build_concept_section.py --out <deck.pptx> [--base <deck.pptx>]

Slides: divider · the chain · the line · the mechanic · five channels · a year · measurement · tests and asks.

Guards, because a concept slide gets challenged in the room:
- every archive title in `checks.archive_titles` must exist in the Sarsabz channel inventory, and every brief
  phrase in `checks.brief_phrases` must exist in the brief text - otherwise it refuses to build;
- exactly one line alternative must be `recommended`;
- `calendar` must reach a month-twelve entry, or the idea is a launch, not a platform.

--base copies an existing deck (never edits it) and appends. Output never overwrites: a clash becomes -v2.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).parent
WS = HERE.parent.parent
NAVY = RGBColor(0x10, 0x2A, 0x43)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x74, 0x80)
ACCENT = RGBColor(0xC8, 0x10, 0x2E)
PALE = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SKY = RGBColor(0x9F, 0xB4, 0xC7)


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("’", "'")).strip().lower()


def validate(c: dict) -> None:
    errors = []
    titles = norm(" || ".join(r["title"] or "" for r in
                              csv.DictReader((WS / "case2_audit/owned/channel_inventory.csv").open(encoding="utf-8"))))
    brief = norm((WS / "brief/full.txt").read_text(encoding="utf-8"))
    for t in c["checks"]["archive_titles"]:
        if norm(t) not in titles:
            errors.append(f"archive title not found in channel inventory: {t!r}")
    for p in c["checks"]["brief_phrases"]:
        if norm(p) not in brief:
            errors.append(f"brief phrase not found in brief text: {p!r}")
    if sum(1 for a in c["line"]["alternatives"] if a.get("recommended")) != 1:
        errors.append("exactly one line alternative must be marked recommended")
    if not any("twelve" in (e.get("stage") or "").lower() for e in c["calendar"]):
        errors.append("calendar has no month-twelve entry - that makes it a launch, not a platform")
    if errors:
        raise SystemExit("Refusing to build:\n  - " + "\n  - ".join(errors))


def txt(slide, l, t, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, space=6, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(line)
        p.alignment = align
        p.space_after = Pt(space)
        f = p.runs[0].font
        f.size, f.bold, f.italic, f.color.rgb, f.name = Pt(size), bold, italic, color, "Segoe UI"
    return tb


def rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def frame(prs, kicker, title, conclusion=None, source=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 1.0, NAVY)
    txt(s, 0.6, 0.12, 12.1, 0.3, kicker, 10, True, SKY, space=0)
    txt(s, 0.6, 0.42, 12.1, 0.5, title, 22, True, WHITE, space=0)
    if conclusion:
        rect(s, 0.6, 6.35, 12.1, 0.62, NAVY)
        txt(s, 0.85, 6.45, 11.6, 0.5, conclusion, 12, True, WHITE, space=0)
    if source:
        txt(s, 0.6, 7.05, 12.1, 0.35, source, 8, False, MUTED, space=0)
    return s


def build(prs, c: dict) -> None:
    ch = c["chain"]
    src = ("Sources: Fatima Creative Marketing Brief 2026; 200 coded owned-YouTube uploads; 1,433 farmer comments; "
           "Sarsabz channel archive (401 uploads). Plot numbers and timings are proposals, not client facts.")

    # divider
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, NAVY)
    rect(s, 0.9, 2.0, 1.6, 0.06, ACCENT)
    txt(s, 0.9, 2.3, 11.4, 0.6, "The campaign", 22, False, SKY)
    txt(s, 0.9, 2.95, 11.6, 1.2, ch["line"], 48, True, WHITE)
    txt(s, 0.9, 4.2, 11.4, 0.5, ch["line_translation"], 20, False, SKY, italic=True)
    txt(s, 0.9, 6.3, 11.4, 0.4, f"Tactical campaign for the 10% promise  ·  from the Big Idea: {c['big_idea']}",
        13, False, SKY)

    # the chain
    s = frame(prs, "THE CAMPAIGN  ·  THE CHAIN", "From the Big Idea to the campaign, one step at a time",
              "Each step earns the next: an asset rivals can't copy, resolving a tension the category avoids, in a space nobody holds.",
              src)
    rows = [("Asset", ch["asset"]), ("Tension", ch["tension"]), ("Territory", ch["territory"]),
            ("Proposition", ch["proposition"]), ("Line", f"{ch['line']}  -  {ch['line_translation']}"),
            ("Mechanic", ch["mechanic"])]
    y = 1.3
    for i, (label, body) in enumerate(rows):
        hi = label == "Line"
        rect(s, 0.6, y, 12.1, 0.78, PALE if hi or i % 2 == 0 else WHITE)
        if hi:
            rect(s, 0.6, y, 0.06, 0.78, ACCENT)
        txt(s, 0.85, y + 0.22, 2.0, 0.4, label.upper(), 11, True, ACCENT if hi else NAVY, space=0)
        txt(s, 2.9, y + 0.12, 9.6, 0.6, body, 13 if hi else 11.5, hi, NAVY if hi else INK, space=0)
        y += 0.82

    # the line
    ln = c["line"]
    s = frame(prs, "THE CAMPAIGN  ·  THE LINE", "Revive the line Sarsabz already owns",
              "Recommended: Dus Feesad Aur. Khet Gawah Hai. - an owned line, with a sign-off that makes the field, not the brand, the one making the claim.",
              "Archive: 'Dus Feesad Aur! | Wheat Crop' (2018, 370K views); 'Dus Feesad Aur | Rice Crop TVC' (2021, 1.8M views); "
              "'Khaad Muft He Samjho | Cotton' (2022, 1.1M views). Script mix from 1,433 farmer comments.")
    rect(s, 0.6, 1.35, 0.06, 1.35, ACCENT)
    txt(s, 0.85, 1.35, 6.2, 0.6, ch["line"], 26, True, NAVY, space=0)
    txt(s, 0.85, 2.15, 6.2, 0.4, ch["line_translation"], 15, False, MUTED, space=0, italic=True)
    txt(s, 0.6, 3.0, 6.4, 0.3, "WHY THIS LINE", 10, True, MUTED, space=0)
    txt(s, 0.6, 3.35, 6.4, 2.9, [f"·  {w}" for w in ln["why"]], 11.5, False, INK, space=9)
    rect(s, 7.4, 1.35, 5.3, 4.85, PALE)
    txt(s, 7.65, 1.5, 4.8, 0.3, "LINES CONSIDERED", 10, True, MUTED, space=0)
    y = 1.9
    for a in ln["alternatives"]:
        rec = a.get("recommended")
        txt(s, 7.65, y, 4.8, 0.3, ("Recommended - " if rec else "") + a["name"], 11, True,
            ACCENT if rec else NAVY, space=0)
        txt(s, 7.65, y + 0.3, 4.8, 0.6, a["tradeoff"], 10, False, INK, space=0)
        y += 0.9
    txt(s, 7.65, 5.45, 4.8, 0.7, f"Swap test: {ln['swap_test']}", 10, True, NAVY, space=0)

    # the mechanic
    s = frame(prs, "THE CAMPAIGN  ·  THE MECHANIC", "Do Khet, Kaanta Din, Nuskha: the idea that refreshes itself",
              "Every result is published - including any under 10%. That is what makes the field a witness rather than an advert.",
              src)
    for i, m in enumerate(c["mechanic"]):
        x = 0.6 + i * 4.13
        rect(s, x, 1.4, 3.85, 0.06, ACCENT)
        txt(s, x, 1.65, 3.85, 0.3, f"STEP {i + 1}", 10, True, MUTED, space=0)
        txt(s, x, 1.95, 3.85, 0.6, m["name"], 26, True, NAVY, space=0)
        txt(s, x, 2.6, 3.85, 0.4, m["gloss"], 13, False, ACCENT, space=0, italic=True)
        txt(s, x, 3.15, 3.85, 2.5, m["body"], 12.5, False, INK, space=0)
    txt(s, 0.6, 5.55, 12.1, 0.6,
        "Repeated every crop, every season, every district - wheat and CAN top-dressing in Rabi, rice, cotton and maize in Kharif.",
        12, True, NAVY, space=0)

    # five channels
    s = frame(prs, "THE CAMPAIGN  ·  HOW IT TRAVELS", "One mechanic across the brief's five channels",
              "Every channel carries the same thing: a real field, both weights, and the recipe.",
              "Channels as listed in the brief (p13). Audience figures: presence sweep 15 Sep 2026. Farmer behaviour: brief p7.")
    txt(s, 0.6, 1.3, 1.9, 0.3, "CHANNEL", 10, True, MUTED, space=0)
    txt(s, 2.6, 1.3, 6.3, 0.3, "EXECUTION", 10, True, MUTED, space=0)
    txt(s, 9.1, 1.3, 3.6, 0.3, "WHAT IT ANSWERS", 10, True, MUTED, space=0)
    y = 1.65
    for i, e in enumerate(c["executions"]):
        if i % 2 == 0:
            rect(s, 0.6, y - 0.05, 12.1, 0.92, PALE)
        txt(s, 0.75, y + 0.1, 1.8, 0.6, e["channel"], 13, True, NAVY, space=0)
        txt(s, 2.6, y + 0.05, 6.3, 0.85, e["work"], 10.5, False, INK, space=0)
        txt(s, 9.1, y + 0.05, 3.5, 0.85, e["answers"], 10, False, MUTED, space=0, italic=True)
        y += 0.92

    # a year
    s = frame(prs, "THE CAMPAIGN  ·  ACROSS A YEAR", "Rabi to Rabi: month one to month twelve",
              "In month twelve the campaign has more proof than it started with. A platform generates work; a launch repeats itself.",
              "Seasons: brief p6 (Rabi Oct - Mar, Kharif Apr - Sep). Timings are a proposal to be set with Fatima agronomy.")
    n = len(c["calendar"])
    w = 12.1 / n
    for i, e in enumerate(c["calendar"]):
        x = 0.6 + i * w
        twelve = "twelve" in e["stage"].lower()
        rect(s, x + 0.04, 1.4, w - 0.08, 0.06, ACCENT if twelve else NAVY)
        txt(s, x + 0.04, 1.6, w - 0.12, 0.5, e["when"], 11, True, ACCENT if twelve else NAVY, space=0)
        txt(s, x + 0.04, 2.1, w - 0.12, 0.4, e["stage"].upper(), 10, True, MUTED, space=0)
        rect(s, x + 0.04, 2.55, w - 0.08, 2.3, PALE if twelve else WHITE)
        txt(s, x + 0.12, 2.65, w - 0.26, 3.4, e["work"], 11, False, INK, space=0)

    # measurement
    s = frame(prs, "THE CAMPAIGN  ·  EFFECTIVENESS", "How we will know it worked",
              "Plot districts against matched districts without plots: the campaign's own design gives Fatima a control group.",
              "The brief asks for 'a practical and measurable approach to campaign effectiveness' (p14). Baselines marked 'requested' "
              "need client data.")
    txt(s, 0.6, 1.3, 1.6, 0.3, "KIND", 10, True, MUTED, space=0)
    txt(s, 2.3, 1.3, 6.6, 0.3, "MEASURE", 10, True, MUTED, space=0)
    txt(s, 9.1, 1.3, 3.6, 0.3, "BASELINE", 10, True, MUTED, space=0)
    y = 1.7
    for i, m in enumerate(c["measurement"]):
        rect(s, 0.6, y - 0.05, 12.1, 1.35, PALE if i % 2 == 0 else WHITE)
        txt(s, 0.75, y + 0.2, 1.5, 0.5, m["kind"], 16, True, NAVY, space=0)
        txt(s, 2.3, y + 0.15, 6.6, 1.1, m["measure"], 12, False, INK, space=0)
        txt(s, 9.1, y + 0.15, 3.5, 1.1, m["baseline"], 11.5, True, ACCENT if "requested" in m["baseline"] else NAVY,
            space=0)
        y += 1.45

    # tests and asks
    s = frame(prs, "THE CAMPAIGN  ·  WHY IT HOLDS", "The six tests, and what the campaign asks of Fatima",
              "If Fatima will not publish results that fall short, do not run 'Khet Gawah Hai' - run 'Dus Feesad Aur' with corrected Ki Jeet stories, and call it a promise, not proof.",
              "Tests: campaign-concept skill (swap, asset, deletion, department, month twelve, room).")
    txt(s, 0.6, 1.3, 6.0, 0.3, "THE SIX TESTS", 10, True, MUTED, space=0)
    y = 1.65
    for t in c["tests"]:
        txt(s, 0.6, y, 1.6, 0.3, t["name"], 11.5, True, NAVY, space=0)
        txt(s, 2.2, y, 4.5, 0.7, t["result"], 10.5, False, INK, space=0)
        y += 0.76
    rect(s, 7.2, 1.3, 5.5, 4.9, PALE)
    txt(s, 7.45, 1.45, 5.0, 0.3, "WHAT IT ASKS OF FATIMA - BEFORE LAUNCH", 10, True, MUTED, space=0)
    y = 1.85
    for i, a in enumerate(c["asks"], 1):
        rect(s, 7.45, y + 0.02, 0.05, 0.7, ACCENT)
        txt(s, 7.65, y, 4.85, 0.8, f"{i}.  {a}", 10.5, False, INK, space=0)
        y += 0.86


def free_path(p: Path) -> Path:
    if not p.exists():
        return p
    n = 2
    while (q := p.with_name(f"{p.stem}-v{n}{p.suffix}")).exists():
        n += 1
    return q


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--base", type=Path)
    ap.add_argument("--spec", type=Path, default=HERE / "concept.yaml")
    args = ap.parse_args()

    c = yaml.safe_load(args.spec.read_text(encoding="utf-8"))
    validate(c)
    if args.base:
        prs = Presentation(str(args.base))
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    before = len(prs.slides)
    build(prs, c)
    out = free_path(args.out)
    prs.save(out)
    print(f"{out}  ({len(prs.slides)} slides; concept section = {len(prs.slides) - before})")


if __name__ == "__main__":
    main()
